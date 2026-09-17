"""
Renderers for communication artifacts.

Four output formats, one input: an artifact whose blocks are already resolved.
No renderer reads a dataset, recomputes a statistic or reformats a value on its
own — they receive finished text and lay it out.  asks for exactly this, and
the reason is fidelity: a number that were re-derived per format could differ
between the DOCX and the slides, and the reader would have no way to know which
was right.

Rendering refuses to proceed on an artifact whose integrity check fails. A
document that cannot be trusted should not exist as a file, because the file
outlives the warning that would have accompanied it on screen.
"""

from __future__ import annotations

import html
import io
from typing import Any

from . import citations as citations_mod
from . import communication
from .ids import new_id
from .storage import export_directory, export_path, storage_root

FORMATS = ("markdown", "html", "docx", "pptx")


class RenderError(RuntimeError):
    """The artifact could not be rendered."""


def render(cur, *, artifact_id: str, fmt: str) -> dict[str, Any]:
    """
    Render an artifact, refusing if its integrity check found problems.

    The refusal carries the problems rather than a generic failure, because
     wants the researcher told what is wrong and what was preserved — and
    because "3 blocks reference an analysis run that no longer exists" is
    actionable in a way that "export failed" is not.
    """
    if fmt not in FORMATS:
        raise RenderError(
            f"Unsupported format {fmt!r}. This installation renders: {', '.join(FORMATS)}."
        )

    integrity = communication.check_integrity(cur, artifact_id)
    if not integrity["publishable"]:
        raise RenderError(
            "This artifact has unresolved problems and will not be rendered: "
            + "; ".join(p["detail"] for p in integrity["problems"])
        )

    artifact = communication.load_artifact(cur, artifact_id, resolve=True)

    if fmt == "markdown":
        payload, suffix = _markdown(artifact).encode("utf-8"), "md"
    elif fmt == "html":
        payload, suffix = _html(artifact).encode("utf-8"), "html"
    elif fmt == "docx":
        payload, suffix = _docx(artifact), "docx"
    else:
        payload, suffix = _pptx(artifact), "pptx"

    # Same layout as figure renders: content under the storage root, keyed by
    # artifact, so one backup covers every produced file.
    #
    # The filename carries the render id, not the artifact id. It used to be
    # `{artifact_id}.{suffix}`, which meant every render of a format overwrote
    # the last one while still inserting a new row — so each earlier row
    # recorded a `byte_size` and `resolved_hash` for bytes that were no longer
    # there, and its `storage_key` resolved to the newest file instead. Nothing
    # served renders by id yet, so this had never produced a wrong download; the
    # first route that did would have handed back a different document than the
    # row described, and it would have looked like a database fault rather than
    # a naming one.
    #
    # Keeping every render costs disk that the old scheme did not. That is the
    # right trade here: `artifact_staleness` exists to answer what a given
    # export said, and an export whose bytes were silently replaced cannot
    # answer it.
    render_id = new_id("ren")
    directory = export_directory("communication_artifacts", artifact_id)
    directory.mkdir(parents=True, exist_ok=True)
    path = export_path(
        "communication_artifacts", artifact_id, render_id, suffix)
    path.write_bytes(payload)
    storage_key = str(path.relative_to(storage_root()))

    cur.execute(
        "INSERT INTO artifact_renders(id, artifact_id, fmt, storage_key, byte_size, "
        "resolved_hash, artifact_version) VALUES (%s, %s, %s, %s, %s, %s, %s)",
        (render_id, artifact_id, fmt, storage_key, len(payload),
         communication.resolved_hash(artifact), artifact["version"]),
    )
    return {
        "render_id": render_id,
        "fmt": fmt,
        "storage_key": storage_key,
        "byte_size": len(payload),
        "warnings": integrity["warnings"],
    }


# ---------------------------------------------------------------------------
# Shared assembly
# ---------------------------------------------------------------------------

def _references(artifact: dict[str, Any]) -> list[tuple[int, dict[str, Any]]]:
    """
    Number the citations across the whole artifact, first appearance first.

    Deduplicated by citation id, so the same source cited in three paragraphs is
    one entry — and the numbering is stable across formats, which is what lets a
    reader move between the DOCX and the slides.
    """
    seen: dict[str, int] = {}
    ordered: list[tuple[int, dict[str, Any]]] = []
    for block in artifact["blocks"]:
        for citation in block["citations"]:
            if citation["id"] not in seen:
                seen[citation["id"]] = len(ordered) + 1
                ordered.append((len(ordered) + 1, citation))
    return ordered


def _marks(artifact: dict[str, Any], block: dict[str, Any]) -> str:
    """Superscript-style markers for the citations on one block."""
    numbering = {c["id"]: n for n, c in _references(artifact)}
    nums = sorted(numbering[c["id"]] for c in block["citations"])
    return f" [{','.join(str(n) for n in nums)}]" if nums else ""


def _provenance_lines(artifact: dict[str, Any]) -> list[str]:
    """
    Where every displayed number came from.

    This is the  traceability trail in text form, and it is not optional
    decoration: it is what makes the exported file auditable once it has left
    the application, which is precisely when this rule is hardest to honour.
    """
    lines: list[str] = []
    for block in artifact["blocks"]:
        for ref in block.get("value_provenance", []):
            lines.append(f"{ref['name']} = {block['resolved'][ref['name']]} "
                         f"— {ref['source']}, {ref['path']}")
    return lines


# ---------------------------------------------------------------------------
# Markdown
# ---------------------------------------------------------------------------

def _markdown(artifact: dict[str, Any]) -> str:
    out: list[str] = [f"# {artifact['title']}", ""]
    if artifact["purpose"]:
        out += [f"*{artifact['purpose']}*", ""]

    for block in artifact["blocks"]:
        kind = block["block_type"]
        text = block["text"]
        mark = _marks(artifact, block)

        if kind in ("heading", "slide_title"):
            out += [f"## {text}", ""]
        elif kind in ("list", "slide_bullet"):
            out += [f"- {text}{mark}"]
        elif kind == "quote":
            out += [f"> {text}{mark}", ""]
        elif kind == "limitation":
            out += [f"**Limitation.** {text}{mark}", ""]
        elif kind == "figure":
            out += [f"**Figure.** {text}{mark}", ""]
        elif kind == "caption":
            out += [f"*{text}{mark}*", ""]
        else:
            out += [f"{text}{mark}", ""]

    references = _references(artifact)
    if references:
        out += ["", "## References", ""]
        for number, citation in references:
            entry = citations_mod.format_reference(citation)
            # The entailment state travels with the reference. A reader deserves
            # to know which citations were actually checked.
            out.append(f"{number}. {entry} — *{citation['entailment']}*")
        out.append("")

    provenance = _provenance_lines(artifact)
    if provenance:
        out += ["## How every number here was produced", ""]
        out += [f"- `{line}`" for line in provenance]
        out += ["",
                "*Each value above was read from the recorded analysis run at render "
                "time. None was transcribed into this document.*", ""]

    return "\n".join(out)


# ---------------------------------------------------------------------------
# HTML
# ---------------------------------------------------------------------------

def _html(artifact: dict[str, Any]) -> str:
    e = html.escape
    parts: list[str] = [
        "<!doctype html><html lang='en'><head><meta charset='utf-8'>",
        f"<title>{e(artifact['title'])}</title>",
        "<style>",
        "body{max-width:46rem;margin:3rem auto;padding:0 1.25rem;",
        "font:16px/1.65 Georgia,'Iowan Old Style',serif;color:#16181d}",
        "h1,h2{font-weight:600;letter-spacing:-.01em;line-height:1.2}",
        "sup{font-size:.7em}",
        ".limitation{border-left:3px solid #c8a45c;padding-left:.9rem;color:#4a4f57}",
        ".refs,.prov{font-size:.82rem;color:#4a4f57;border-top:1px solid #e4e2dd;",
        "margin-top:2.5rem;padding-top:1rem}",
        ".prov code{font:12px ui-monospace,SFMono-Regular,Menlo,monospace}",
        ".state{font-style:italic;color:#6b7079}",
        "</style></head><body>",
        f"<h1>{e(artifact['title'])}</h1>",
    ]
    if artifact["purpose"]:
        parts.append(f"<p><em>{e(artifact['purpose'])}</em></p>")

    for block in artifact["blocks"]:
        kind = block["block_type"]
        text = e(block["text"])
        mark = _marks(artifact, block)
        sup = f"<sup>{e(mark.strip())}</sup>" if mark else ""

        if kind in ("heading", "slide_title"):
            parts.append(f"<h2>{text}</h2>")
        elif kind in ("list", "slide_bullet"):
            parts.append(f"<ul><li>{text}{sup}</li></ul>")
        elif kind == "quote":
            parts.append(f"<blockquote>{text}{sup}</blockquote>")
        elif kind == "limitation":
            parts.append(f"<p class='limitation'>{text}{sup}</p>")
        elif kind == "caption":
            parts.append(f"<p><em>{text}{sup}</em></p>")
        else:
            parts.append(f"<p>{text}{sup}</p>")

    references = _references(artifact)
    if references:
        parts.append("<div class='refs'><h2>References</h2><ol>")
        for _, citation in references:
            parts.append(
                f"<li>{e(citations_mod.format_reference(citation))} "
                f"<span class='state'>— {e(citation['entailment'])}</span></li>"
            )
        parts.append("</ol></div>")

    provenance = _provenance_lines(artifact)
    if provenance:
        parts.append("<div class='prov'><h2>How every number here was produced</h2><ul>")
        parts += [f"<li><code>{e(line)}</code></li>" for line in provenance]
        parts.append("</ul><p>Each value was read from the recorded analysis run at "
                     "render time. None was transcribed into this document.</p></div>")

    parts.append("</body></html>")
    return "".join(parts)


# ---------------------------------------------------------------------------
# DOCX
# ---------------------------------------------------------------------------

def _docx(artifact: dict[str, Any]) -> bytes:
    from docx import Document
    from docx.shared import Pt

    document = Document()
    document.add_heading(artifact["title"], level=0)
    if artifact["purpose"]:
        document.add_paragraph(artifact["purpose"]).italic = True

    for block in artifact["blocks"]:
        kind = block["block_type"]
        text = block["text"] + _marks(artifact, block)

        if kind in ("heading", "slide_title"):
            document.add_heading(block["text"], level=1)
        elif kind in ("list", "slide_bullet"):
            document.add_paragraph(text, style="List Bullet")
        elif kind == "quote":
            document.add_paragraph(text, style="Intense Quote")
        elif kind == "limitation":
            paragraph = document.add_paragraph()
            run = paragraph.add_run("Limitation. ")
            run.bold = True
            paragraph.add_run(text)
        elif kind == "caption":
            paragraph = document.add_paragraph(text)
            paragraph.runs[0].italic = True
            paragraph.runs[0].font.size = Pt(9)
        else:
            document.add_paragraph(text)

    references = _references(artifact)
    if references:
        document.add_heading("References", level=1)
        for number, citation in references:
            document.add_paragraph(
                f"{number}. {citations_mod.format_reference(citation)} "
                f"— {citation['entailment']}"
            )

    provenance = _provenance_lines(artifact)
    if provenance:
        document.add_heading("How every number here was produced", level=1)
        for line in provenance:
            document.add_paragraph(line, style="List Bullet")
        document.add_paragraph(
            "Each value above was read from the recorded analysis run at render time. "
            "None was transcribed into this document."
        )

    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


# ---------------------------------------------------------------------------
# PPTX
# ---------------------------------------------------------------------------

def _pptx(artifact: dict[str, Any]) -> bytes:
    """
    One idea per slide.

    Blocks are grouped so a `heading` or `slide_title` opens a slide and the
    following blocks become its bullets. the system forbids pasting manuscript
    paragraphs onto slides, so prose is truncated with an ellipsis and the full
    text goes to the speaker notes, where it belongs.
    """
    from pptx import Presentation
    from pptx.util import Inches, Pt

    presentation = Presentation()
    presentation.slide_width = Inches(13.333)   # 16:9
    presentation.slide_height = Inches(7.5)

    title_layout = presentation.slide_layouts[0]
    bullet_layout = presentation.slide_layouts[1]

    opening = presentation.slides.add_slide(title_layout)
    opening.shapes.title.text = artifact["title"]
    if len(opening.placeholders) > 1:
        opening.placeholders[1].text = artifact["purpose"] or "Throughline"

    current = None
    body = None

    def new_slide(heading: str):
        slide = presentation.slides.add_slide(bullet_layout)
        slide.shapes.title.text = heading
        frame = slide.placeholders[1].text_frame
        frame.clear()
        return slide, frame

    notes: list[str] = []

    for block in artifact["blocks"]:
        kind = block["block_type"]
        text = block["text"] + _marks(artifact, block)

        if kind in ("heading", "slide_title") or current is None:
            if current is not None and notes:
                current.notes_slide.notes_text_frame.text = "\n\n".join(notes)
            notes = []
            heading = block["text"] if kind in ("heading", "slide_title") else artifact["title"]
            current, body = new_slide(heading)
            if kind in ("heading", "slide_title"):
                continue

        if block["notes"]:
            notes.append(block["notes"])

        #  — a slide carries a message, not a paragraph.
        shown = text if len(text) <= 220 else text[:217].rstrip() + "…"
        if len(text) > 220:
            notes.append(text)

        paragraph = body.paragraphs[0] if not body.paragraphs[0].text else body.add_paragraph()
        paragraph.text = shown
        paragraph.font.size = Pt(18)
        if kind == "limitation":
            paragraph.level = 1

    if current is not None and notes:
        current.notes_slide.notes_text_frame.text = "\n\n".join(notes)

    references = _references(artifact)
    if references:
        _, frame = new_slide("References")
        for number, citation in references:
            paragraph = frame.paragraphs[0] if not frame.paragraphs[0].text else frame.add_paragraph()
            paragraph.text = (f"{number}. {citations_mod.format_reference(citation)} "
                              f"— {citation['entailment']}")
            paragraph.font.size = Pt(12)

    provenance = _provenance_lines(artifact)
    if provenance:
        _, frame = new_slide("How every number here was produced")
        for line in provenance:
            paragraph = frame.paragraphs[0] if not frame.paragraphs[0].text else frame.add_paragraph()
            paragraph.text = line
            paragraph.font.size = Pt(11)

    buffer = io.BytesIO()
    presentation.save(buffer)
    return buffer.getvalue()


__all__ = ["FORMATS", "RenderError", "render"]
